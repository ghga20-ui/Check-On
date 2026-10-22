"""Import-compatible sample workbooks containing synthetic students only."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def save_template(path, mixed=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "학생 명단"
    sheet.append(["반", "번호", "이름"] if mixed else ["번호", "이름"])
    for i in range(1, 4):
        sheet.append([1 if i < 3 else 2, i, f"예시학생{i}"] if mixed else [i, f"예시학생{i}"])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="176AE6")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for column in ("A", "B", "C"):
        sheet.column_dimensions[column].width = 22
    guide = workbook.create_sheet("작성 안내")
    for line in ["첫 번째 ‘학생 명단’ 시트에 입력한 뒤 저장해 주세요.",
                 "예시학생 행은 지우고 실제 명단으로 바꿔 주세요.",
                 "첫 행의 열 이름을 유지해 주세요. 번호는 1~99, 이름은 생략 가능합니다.",
                 "혼합학급은 원소속 반과 번호를 입력합니다. 같은 반·번호를 중복 입력하지 마세요." if mixed else "선택한 학급의 학생만 입력합니다. 번호를 중복 입력하지 마세요.",
                 "학생 이름은 현재 PC에만 보관되며 Drive로 보내지 않습니다."]:
        guide.append([line])
    guide.column_dimensions["A"].width = 105
    workbook.save(path)
    workbook.close()
